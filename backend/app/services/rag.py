"""Recuperación aumentada (RAG) del material de estudio (HU-003).

LangChain orquesta el proceso y ChromaDB es la base vectorial (Entrega 3, Estructura de Datos).
El procesamiento tiene cinco etapas; aquí se implementan las cuatro primeras (el contraste entre
fuentes usa el modelo de lenguaje):

1. Extracción: PDF página por página, DOCX sección por sección y PPTX diapositiva por diapositiva,
   con la página, sección o diapositiva de origen en `pagina_o_seccion`. Antes de fragmentar, el texto
   se normaliza (ligaduras, glifos ilegibles y, en los PDF, palabras cortadas con guion).
2. Segmentación: fragmentos de hasta 400 caracteres con 60 de superposición.
3. Representación vectorial: paraphrase-multilingual-MiniLM-L12-v2, local (FastEmbed, sin costo por uso).
4. Recuperación: los seis mejores fragmentos de cada fuente (propia y académica), filtrados por tema:
   de los candidatos más similares se descartan los que no alcanzan la similitud coseno mínima, un
   cross-encoder multilingüe los reordena según si responden la consulta (re-ranking) y se eligen
   con un máximo por origen (los lugares libres se completan con más candidatos de otras fuentes).
   A cada fragmento se le agrega su
   contexto (los fragmentos vecinos de la misma sección), y los vecinos recuperados se fusionan.

Cada fragmento guarda los campos de FRAGMENTO_MATERIAL (Figura 6): id_fragmento, id_origen, id_tema,
texto_original, embedding, fuente, nombre_archivo y pagina_o_seccion. `id_origen` es el id del
material propio o de la fuente académica, según el valor de `fuente`. Los valores de cada etapa
están en parametros.toml.
"""
import io
import re
import threading
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from ..config import P, settings

R = P["rag"]
COLECCION = "fragmento_material"
ACADEMICA, PROPIA = "académica", "propia"
TAM_FRAGMENTO = R["tam_fragmento"]
SOLAPAMIENTO = R["solapamiento"]


@dataclass
class Seccion:
    texto: str
    pagina_o_seccion: str  # "pág. 3", "diap. 2", el título de la sección o "" si el archivo no tiene partes


# ---------------- Validación del tipo por contenido ----------------

def tipo_real(contenido: bytes) -> str | None:
    """Tipo del archivo según su contenido, no su extensión (Seguridad: medidas complementarias).

    PDF empieza con "%PDF-". DOCX y PPTX son paquetes ZIP de Office: se distinguen por la carpeta
    principal que declaran (word/ o ppt/). Cualquier otra cosa devuelve None.
    """
    if contenido.startswith(b"%PDF-"):
        return "pdf"
    if not contenido.startswith(b"PK\x03\x04"):
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            nombres = set(z.namelist())
    except zipfile.BadZipFile:
        return None
    if "[Content_Types].xml" not in nombres:
        return None
    if "word/document.xml" in nombres:
        return "docx"
    if "ppt/presentation.xml" in nombres:
        return "pptx"
    return None


# ---------------- Extracción de texto ----------------

def _titulo(linea: str) -> str | None:
    """Títulos del texto de las fuentes (== Título ==).

    No se reconoce "# Título" de Markdown: en los ejemplos de código un comentario de Python
    empieza igual y se confundía con un título de sección.
    """
    m = re.match(r"^\s*={2,}\s*(.+?)\s*={2,}\s*$", linea)
    return m.group(1).strip() if m else None


def por_titulos(texto: str) -> list[Seccion]:
    secciones, actual, lineas = [], "", []
    for linea in texto.splitlines():
        t = _titulo(linea)
        if t is None:
            lineas.append(linea)
            continue
        secciones.append(Seccion("\n".join(lineas), actual))
        actual, lineas = t, []
    secciones.append(Seccion("\n".join(lineas), actual))
    return secciones


def _pptx(ruta: Path) -> list[Seccion]:
    from pptx import Presentation
    secciones = []
    for n, diapositiva in enumerate(Presentation(ruta).slides, 1):
        partes = []
        for forma in diapositiva.shapes:
            if forma.has_text_frame:
                partes.append(forma.text_frame.text)
            elif getattr(forma, "has_table", False) and forma.has_table:
                partes += [" | ".join(c.text for c in fila.cells) for fila in forma.table.rows]
        if diapositiva.has_notes_slide:  # las notas del orador suelen tener la explicación
            partes.append(diapositiva.notes_slide.notes_text_frame.text)
        secciones.append(Seccion("\n".join(p for p in partes if p.strip()), f"diap. {n}"))
    return secciones


def extraer(ruta: Path, tipo: str) -> list[Seccion]:
    if tipo == "pdf":
        from pypdf import PdfReader
        return [Seccion(unir_guiones(p.extract_text() or ""), f"pág. {i}") for i, p in enumerate(PdfReader(ruta).pages, 1)]
    if tipo == "docx":
        from docx import Document
        secciones, actual, parrafos = [], "", []
        for p in Document(ruta).paragraphs:
            if p.style.name.lower().startswith(("heading", "título")) and p.text.strip():
                secciones.append(Seccion("\n".join(parrafos), actual))
                actual, parrafos = p.text.strip(), []
            else:
                parrafos.append(p.text)
        secciones.append(Seccion("\n".join(parrafos), actual))
        return secciones
    if tipo == "pptx":
        return _pptx(ruta)
    # Texto plano con títulos: lo usan las fuentes académicas precargadas (no se admite como material propio).
    return por_titulos(ruta.read_text(encoding="utf-8", errors="replace"))


# ---------------- Normalización ----------------

# Cortes de palabra con guion al final de una línea de PDF ("recur-\nsive"): solo si sigue en minúscula,
# porque "Floyd-\nPratt" es un nombre compuesto.
CORTE_CON_GUION = re.compile(r"([^\W\d_])-[ \t]*\n[ \t]*(?=[^\W\d_A-ZÁÉÍÓÚÑ])")


def unir_guiones(texto: str) -> str:
    """Une las palabras que el PDF cortó con guion al final de la línea (se aplica solo a PDF)."""
    return CORTE_CON_GUION.sub(r"\1", texto)


def _ilegible(c: str) -> bool:
    """Glifos sin equivalente Unicode: los PDF con tipografías propias los extraen como U+FFFF o U+FFFD, y
    las versalitas de Erickson, como caracteres de control (U+0000). Tabuladores y saltos de línea no cuentan."""
    return c == "�" or unicodedata.category(c) in ("Cn", "Co") or (unicodedata.category(c) == "Cc" and c not in "\t\r\n")


def normalizar(texto: str) -> str:
    """Texto listo para fragmentar, de cualquier formato (Estructura de Datos, etapa de extracción).

    - NFKC: las ligaduras tipográficas pasan a letras ("ﬁ" → "fi") y los espacios especiales, a espacios.
      Sin esto, "ﬁnite" y "finite" son palabras distintas para el modelo de embeddings.
    - Los glifos ilegibles se quitan; si son la mayor parte de una línea (encabezados de página de Erickson,
      "￿. R￿￿￿￿￿￿￿￿"), la línea entera se descarta, porque lo que queda no se puede leer.
    """
    texto = unicodedata.normalize("NFKC", texto)
    lineas = []
    for linea in texto.split("\n"):
        visibles = [c for c in linea if not c.isspace()]
        malos = sum(_ilegible(c) for c in visibles)
        if malos and malos / len(visibles) >= R["max_ilegible"]:
            continue
        lineas.append("".join(c for c in linea if not _ilegible(c)) if malos else linea)
    return "\n".join(lineas)


# ---------------- Fragmentación ----------------

def fragmentar(texto: str, tam: int = TAM_FRAGMENTO, solape: int = SOLAPAMIENTO) -> list[str]:
    """Corta primero por párrafos, después por líneas, oraciones y palabras, hasta que cada fragmento entre en `tam`.

    El separador de oraciones queda al final del fragmento (keep_separator="end"): así ninguno empieza con
    el punto de la oración anterior.
    """
    divisor = RecursiveCharacterTextSplitter(chunk_size=tam, chunk_overlap=solape, keep_separator="end",
                                             separators=["\n\n", "\n", ". ", " ", ""])
    return [f.strip() for f in divisor.split_text(texto) if f.strip()]


# ---------------- Embeddings y ChromaDB ----------------

class _FastEmbed:
    def __init__(self):
        from fastembed import TextEmbedding
        self.modelo = TextEmbedding(R["modelo_embeddings"])  # se descarga la primera vez (~220 MB)

    def __call__(self, textos: list[str]) -> list[list[float]]:
        return [v.tolist() for v in self.modelo.embed(textos)]


embedder = None  # se crea al primer uso; los tests lo reemplazan por uno falso
_almacen = None
_lock = threading.Lock()  # dos archivos subidos juntos se procesan en hilos distintos: se crean una sola vez


def _embed(textos: list[str]) -> list[list[float]]:
    global embedder
    with _lock:
        if embedder is None:
            embedder = _FastEmbed()
    return embedder(textos)


class _CrossEncoder:
    """Modelo de re-ranking: lee la consulta y cada fragmento juntos y devuelve un puntaje de pertinencia.

    mmarco-mMiniLMv2 es multilingüe (consultas en español, fuentes en inglés) y no viene en la lista de
    FastEmbed: se registra con su archivo ONNX. Se descarga la primera vez (~470 MB).
    """
    def __init__(self):
        from fastembed.common.model_description import ModelSource
        from fastembed.rerank.cross_encoder import TextCrossEncoder
        nombre = R["modelo_reranking"]
        if nombre not in {m["model"] for m in TextCrossEncoder.list_supported_models()}:
            TextCrossEncoder.add_custom_model(model=nombre, sources=ModelSource(hf=nombre), model_file="onnx/model.onnx",
                                              license="apache-2.0", size_in_gb=0.47)
        self.modelo = TextCrossEncoder(nombre)

    def __call__(self, consulta: str, textos: list[str]) -> list[float]:
        return [float(p) for p in self.modelo.rerank(consulta, textos)]


reordenador = None  # se crea al primer uso; los tests lo reemplazan por uno falso


def _puntajes(consulta: str, textos: list[str]) -> list[float]:
    global reordenador
    with _lock:
        if reordenador is None:
            reordenador = _CrossEncoder()
    return reordenador(consulta, textos) if textos else []


class EmbeddingsLocales(Embeddings):
    """Adapta el modelo local a la interfaz de embeddings de LangChain."""

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return _embed(texts)

    def embed_query(self, text: str) -> list[float]:
        return _embed([text])[0]


def almacen():
    """Vector store de LangChain sobre una colección persistente de ChromaDB, con distancia coseno."""
    global _almacen
    with _lock:
        if _almacen is None:
            import chromadb
            from langchain_chroma import Chroma
            _almacen = Chroma(collection_name=COLECCION, embedding_function=EmbeddingsLocales(),
                              client=chromadb.PersistentClient(path=settings.chroma_dir),
                              collection_configuration={"hnsw": {"space": "cosine"}})
    return _almacen


def vaciar() -> None:
    """Borra todos los fragmentos (al reiniciar una base local de prueba)."""
    col = almacen()
    if ids := col.get()["ids"]:
        col.delete(ids=ids)


def indexar(secciones: list[Seccion], *, id_origen: int, id_tema: int, fuente: str, nombre_archivo: str,
            prefijo: str) -> int:
    """Guarda los fragmentos en ChromaDB. Devuelve cuántos guardó (0 si el archivo no tenía texto).

    Los fragmentos muy cortos (un título o una frase suelta) no aportan contenido y, al ser cortos, se
    parecen a casi cualquier consulta: se descartan, salvo que sean todo lo que tiene el archivo.
    """
    docs, metas = [], []
    for s in secciones:
        for f in fragmentar(normalizar(s.texto)):
            docs.append(f)
            metas.append({"id_origen": id_origen, "id_tema": id_tema, "fuente": fuente,
                          "nombre_archivo": nombre_archivo, "pagina_o_seccion": s.pagina_o_seccion})
    if any(len(d) >= R["fragmento_minimo"] for d in docs):
        pares = [(d, m) for d, m in zip(docs, metas) if len(d) >= R["fragmento_minimo"]]
        docs, metas = [d for d, _ in pares], [m for _, m in pares]
    # Posición de cada fragmento dentro de su origen: permite traer los vecinos al recuperar (ver contexto).
    for n, m in enumerate(metas):
        m["orden"] = n
    vs = almacen()
    vs.delete(where={"$and": [{"fuente": fuente}, {"id_origen": id_origen}]})  # reindexar no duplica
    if docs:
        vs.add_texts(docs, metadatas=metas, ids=[f"{prefijo}-{n}" for n in range(len(docs))])  # id_fragmento
    return len(docs)


def _de_una_fuente(consulta: str, filtro: dict, k: int, minima: float, max_por_origen: int) -> list[dict]:
    """Los k mejores fragmentos que superan el umbral, con como mucho `max_por_origen` de cada origen.

    Se piden más candidatos que los que se devuelven, por dos motivos:
    - Re-ranking: la similitud de los embeddings compara la consulta y el fragmento por separado, y los
      fragmentos generales sobre el tema (un glosario, una introducción) se parecen a casi cualquier
      consulta. Un cross-encoder lee la consulta y cada candidato juntos y los reordena según si lo
      responden (puntaje_reranking). El umbral se sigue aplicando a la similitud coseno: el puntaje del
      cross-encoder está en otra escala, y así "suficiente" (CA5) no depende del reordenamiento.
    - Diversidad: si una fuente grande (Algorithms de Erickson tiene unos 850 fragmentos de Recursividad)
      ocupa los primeros puestos, los lugares que deja libres se completan con los siguientes candidatos
      de otras fuentes, siempre sobre el umbral. 0 = sin máximo por origen.
    """
    resultados = almacen().similarity_search_with_relevance_scores(consulta, k=max(R["candidatos"], k), filter=filtro)
    candidatos = [{"id_fragmento": doc.id, "texto_original": doc.page_content, "similitud": round(score, 3),
                   **doc.metadata} for doc, score in resultados if score >= minima]  # de mayor a menor similitud
    if R["reranking"]:
        for f, p in zip(candidatos, _puntajes(consulta, [f["texto_original"] for f in candidatos])):
            f["puntaje_reranking"] = round(p, 3)
        candidatos.sort(key=lambda f: f["puntaje_reranking"], reverse=True)
    elegidos, por_origen = [], {}
    for f in candidatos:
        o = f["id_origen"]
        if max_por_origen and por_origen.get(o, 0) >= max_por_origen:
            continue
        por_origen[o] = por_origen.get(o, 0) + 1
        elegidos.append(f)
        if len(elegidos) == k:
            break
    return elegidos


def _orden(f: dict) -> float:
    """Criterio para ordenar los resultados: el puntaje del re-ranking si se aplicó, si no la similitud."""
    return f.get("puntaje_reranking", f["similitud"])


def buscar(id_tema: int, consulta: str, materiales_propios: list[int], k: int | None = None,
           minima: float | None = None) -> list[dict]:
    """Los k mejores fragmentos de cada fuente del tema, sin los que no alcanzan la similitud mínima.

    Solo fuentes académicas o material del propio estudiante: nunca el material que subió otra persona.
    En cada fuente se limita cuántos fragmentos puede aportar un mismo origen (max_por_origen y
    max_por_origen_propio en parametros.toml), para que un libro grande no acapare los resultados.
    Los parámetros se leen en cada llamada, así la evaluación puede probar otros valores.
    """
    k = R["k_por_fuente"] if k is None else k
    minima = R["similitud_minima"] if minima is None else minima
    encontrados = _de_una_fuente(consulta, {"$and": [{"id_tema": id_tema}, {"fuente": ACADEMICA}]}, k, minima,
                                 R["max_por_origen"])
    if materiales_propios:
        encontrados += _de_una_fuente(consulta, {"$and": [{"id_tema": id_tema}, {"fuente": PROPIA},
                                                          {"id_origen": {"$in": materiales_propios}}]}, k, minima,
                                      R["max_por_origen_propio"])
    # Con re-ranking, los fragmentos propios y los académicos se ordenan juntos por el mismo modelo.
    return con_contexto(sorted(encontrados, key=_orden, reverse=True))


# ---------------- Contexto: fragmentos vecinos ----------------

def unir_solapados(a: str, b: str) -> str:
    """Une dos fragmentos consecutivos sin repetir la superposición (hasta 60 caracteres, ver fragmentar).

    El final de `a` reaparece al principio de `b`: se busca la coincidencia más larga que empiece y
    termine en un límite de palabra, para no confundir una superposición con una palabra repetida por azar.
    """
    for k in range(min(len(a), len(b)), 0, -1):
        if (a.endswith(b[:k]) and (k == len(b) or not b[k].isalnum())
                and (k == len(a) or not a[-k - 1].isalnum())):
            return a + b[k:]
    return f"{a}\n{b}"


def _vecinos(f: dict, desde: int, hasta: int) -> list[str]:
    """Textos de los fragmentos del mismo origen y la misma sección con orden entre `desde` y `hasta`."""
    r = almacen().get(where={"$and": [{"fuente": f["fuente"]}, {"id_origen": f["id_origen"]},
                                      {"pagina_o_seccion": f["pagina_o_seccion"]},
                                      {"orden": {"$gte": desde}}, {"orden": {"$lte": hasta}}]})
    return [t for _, t in sorted(zip((m["orden"] for m in r["metadatas"]), r["documents"]))]


def con_contexto(fragmentos: list[dict], vecinos: int = R["vecinos"]) -> list[dict]:
    """Agrega a cada fragmento recuperado el campo "contexto": el texto con sus vecinos (orden ± vecinos).

    Los fragmentos de 400 caracteres son buenos para comparar con la consulta (cada uno entra completo en
    el vector), pero a veces les falta el párrafo anterior o el siguiente para que se entiendan. La
    búsqueda sigue comparando contra el fragmento, y el contexto se arma después, para quien genera el
    contenido. Solo se juntan fragmentos del mismo origen y la misma página o sección.

    Si las ventanas de dos fragmentos recuperados se tocan o se pisan, se fusionan en un resultado, para
    no repetir texto: queda en la posición del mejor ubicado, con su texto_original y su similitud, y
    `recuperados` lista todos los fragmentos que superaron el umbral dentro de ese contexto.
    """
    grupos: list[dict] = []
    for f in fragmentos:  # vienen ordenados (re-ranking o similitud): el primero de cada grupo es el mejor
        if "orden" not in f:  # indexado antes de guardar el orden: sin vecinos
            grupos.append({**f, "contexto": f["texto_original"], "recuperados": [f["id_fragmento"]]})
            continue
        clave = (f["fuente"], f["id_origen"], f["pagina_o_seccion"])
        g = next((g for g in grupos if g.get("_clave") == clave
                  and f["orden"] - vecinos <= g["_hasta"] + 1 and f["orden"] + vecinos >= g["_desde"] - 1), None)
        if g:
            g["_desde"], g["_hasta"] = min(g["_desde"], f["orden"] - vecinos), max(g["_hasta"], f["orden"] + vecinos)
            g["recuperados"].append(f["id_fragmento"])
        else:
            grupos.append({**f, "_clave": clave, "_desde": f["orden"] - vecinos, "_hasta": f["orden"] + vecinos,
                           "recuperados": [f["id_fragmento"]]})
    for g in grupos:
        if "_clave" in g:
            textos = _vecinos(g, g.pop("_desde"), g.pop("_hasta"))
            del g["_clave"]
            contexto = textos[0] if textos else g["texto_original"]
            for t in textos[1:]:
                contexto = unir_solapados(contexto, t)
            g["contexto"] = contexto
    return grupos


def recuperar(id_tema: int, consulta: str, materiales_propios: list[int]) -> dict:
    """Fragmentos para generar contenido sobre un concepto, e indicación de si alcanzan (HU-003, CA5).

    Con menos de dos fragmentos sobre el umbral, el material es insuficiente: quien llama informa al
    estudiante y solo genera una respuesta del conocimiento general del modelo si él lo pide, marcada
    como no respaldada por las fuentes del tema. Se cuentan los fragmentos que superaron el umbral,
    aunque dos vecinos se hayan fusionado en un mismo resultado.
    """
    fragmentos = buscar(id_tema, consulta, materiales_propios)
    recuperados = sum(len(f["recuperados"]) for f in fragmentos)
    return {"fragmentos": fragmentos, "suficiente": recuperados >= R["fragmentos_minimos"]}
