"""Archivos de prueba reales (PDF, DOCX y PPTX) para la carga de material."""
import io


def pdf(*paginas: str) -> bytes:
    """PDF mínimo válido, una página por texto (sin depender de librerías para generarlo)."""
    objetos = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for texto in paginas:
        contenido = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
        objetos.append(f"<< /Length {len(contenido)} >>\nstream\n{contenido.decode('latin-1')}\nendstream")
        objetos.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objetos)} 0 R "
                       "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objetos)} 0 R")
    objetos[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    salida, offsets = io.BytesIO(), []
    salida.write(b"%PDF-1.4\n")
    for i, o in enumerate(objetos, 1):
        offsets.append(salida.tell())
        salida.write(f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1"))
    xref = salida.tell()
    salida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        salida.write(f"{off:010d} 00000 n \n".encode())
    salida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return salida.getvalue()


def docx(*parrafos: str, titulo: str | None = None) -> bytes:
    from docx import Document
    d = Document()
    if titulo:
        d.add_heading(titulo, level=1)
    for p in parrafos:
        d.add_paragraph(p)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def pptx(*diapositivas: str) -> bytes:
    """Una diapositiva por texto; una cadena vacía da una diapositiva sin texto (como una de solo imágenes)."""
    from pptx import Presentation
    from pptx.util import Inches
    pres = Presentation()
    for texto in diapositivas:
        d = pres.slides.add_slide(pres.slide_layouts[6])  # en blanco
        if texto:
            d.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.text = texto
    b = io.BytesIO(); pres.save(b)
    return b.getvalue()
