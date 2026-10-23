"""RAG del material (HU-003, tareas 06 y 07).

Usa el modelo de embeddings falso de conftest.py: alcanza para comprobar extracción,
fragmentación, metadatos, filtros y orden por similitud.
"""
import io
from datetime import date, timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session, SQLModel, select

from app.db import engine
from app.main import app
from app.models import MaterialEstudio, Tema
from app.services import rag


@pytest.fixture()
def client():
    from app.db import init_db
    init_db()
    SQLModel.metadata.drop_all(engine)
    with TestClient(app) as c:
        yield c


def crear_plan(c, topic="Recursividad"):
    start = date(2026, 9, 21)
    body = dict(topic=topic, start_date=str(start), end_date=str(start + timedelta(days=27)), weekdays=[0, 2, 4],
                level="intermedio", target_grade=8, language="java")
    return c.post("/api/plans", json=body).json()["id"]


def id_tema(nombre):
    with Session(engine) as db:
        return db.exec(select(Tema.id_tema).where(Tema.nombre == nombre)).one()


# ---------------- Fragmentación y extracción ----------------

def test_fragmentos_respetan_tamano_y_solapan():
    texto = "\n\n".join(f"Oración número {i} sobre recursividad y caso base." for i in range(60))
    frags = rag.fragmentar(texto, tam=300, solape=60)
    assert len(frags) > 5 and all(len(f) <= 300 for f in frags)
    assert all(a[-20:].split()[-1] in b for a, b in zip(frags, frags[1:]))  # el final de uno reaparece en el siguiente


def test_las_oraciones_cortadas_no_empiezan_con_punto():
    frags = rag.fragmentar("El caso base corta la recursión y evita llamadas infinitas. " * 30)
    assert all(not f.startswith(".") for f in frags) and all(f.endswith(".") for f in frags)


def test_fragmentos_muy_cortos_no_se_indexan(client):
    tema = id_tema("Recursividad")
    n = rag.indexar([rag.Seccion("Según lo observado arriba:", "A"), rag.Seccion("La recursión resuelve un problema "
                     "llamándose a sí misma sobre una versión más chica del mismo problema.", "B")],
                    id_material=0, id_tema=tema, fuente=rag.ACADEMICA, nombre_archivo="Recursión", prefijo="w5")
    assert n == 1
    # Pero un apunte de una sola línea se guarda igual.
    assert rag.indexar([rag.Seccion("Caso base: corta la recursión.", "")], id_material=7, id_tema=tema,
                       fuente=rag.PROPIA, nombre_archivo="nota.txt", prefijo="mat7") == 1


def test_texto_sin_espacios_se_corta_igual():
    frags = rag.fragmentar("x" * 2000, tam=800, solape=150)
    assert all(len(f) <= 800 for f in frags) and len("".join(frags)) >= 2000


def test_markdown_y_wikipedia_se_dividen_por_titulos():
    secs = rag.por_titulos("Intro\n# Caso base\ntexto A\n== Pila de llamadas ==\ntexto B")
    assert [(s.pagina_o_seccion, s.texto.strip()) for s in secs] == [("", "Intro"), ("Caso base", "texto A"),
                                                                     ("Pila de llamadas", "texto B")]


def test_docx_usa_los_titulos_como_seccion(tmp_path):
    d = Document()
    d.add_paragraph("Introducción general")
    d.add_heading("Caso base", level=1)
    d.add_paragraph("Condición que corta la recursión.")
    d.save(tmp_path / "a.docx")
    secs = rag.extraer(tmp_path / "a.docx", "docx")
    assert ("Caso base", "Condición que corta la recursión.") in [(s.pagina_o_seccion, s.texto) for s in secs]


# ---------------- Flujo completo ----------------

def subir(c, plan_id, nombre, contenido: bytes):
    return c.post(f"/api/plans/{plan_id}/materials", files={"file": (nombre, io.BytesIO(contenido))})


def test_subir_material_lo_indexa_con_los_metadatos_del_der(client):
    client.post("/api/auth/dev")
    plan = crear_plan(client)
    r = subir(client, plan, "apunte.md", "# Caso base\nToda función recursiva necesita un caso base.".encode())
    assert r.status_code == 201
    mats = client.get("/api/materials").json()
    assert mats[0]["status"] == "procesado"
    meta = rag.almacen().get(where={"id_material": mats[0]["id"]})["metadatas"][0]
    assert meta == {"id_material": mats[0]["id"], "id_tema": id_tema("Recursividad"), "fuente": "propia",
                    "nombre_archivo": "apunte.md", "pagina_o_seccion": "Caso base"}


def test_archivo_sin_texto_queda_con_error(client):
    client.post("/api/auth/dev")
    plan = crear_plan(client)
    subir(client, plan, "vacio.txt", b"   ")
    subir(client, plan, "roto.pdf", b"esto no es un pdf")
    assert {m["status"] for m in client.get("/api/materials").json()} == {"con error"}


def test_busqueda_devuelve_lo_mas_parecido_primero(client):
    client.post("/api/auth/dev")
    plan = crear_plan(client)
    tema = id_tema("Recursividad")
    rag.indexar([rag.Seccion("La pila de llamadas guarda cada invocación pendiente de la función.", "Pila")],
                id_material=0, id_tema=tema, fuente=rag.ACADEMICA, nombre_archivo="Pila de llamadas", prefijo="w1")
    rag.indexar([rag.Seccion("El caso base detiene la recursión y evita llamadas infinitas.", "Caso base")],
                id_material=0, id_tema=tema, fuente=rag.ACADEMICA, nombre_archivo="Recursión", prefijo="w2")
    r = client.get(f"/api/plans/{plan}/search", params={"q": "qué es el caso base de la recursión"}).json()
    assert r[0]["nombre_archivo"] == "Recursión" and r[0]["similitud"] > r[1]["similitud"]


def test_busqueda_no_mezcla_temas_ni_material_ajeno(client):
    # Otro estudiante sube material de Recursividad.
    client.post("/api/auth/dev")
    otro = client.get("/api/auth/me").json()["id"]
    plan_otro = crear_plan(client)
    subir(client, plan_otro, "ajeno.txt", "caso base secreto del otro estudiante".encode())
    with Session(engine) as db:
        from app.models import Estudiante
        db.add(Estudiante(id_google="otra-persona", nombre="Ana", correo_electronico="ana@example.com"))
        db.commit()
    # Una fuente académica de otro tema.
    rag.indexar([rag.Seccion("caso base de la búsqueda binaria", "")], id_material=0,
                id_tema=id_tema("Búsqueda y ordenamiento"), fuente=rag.ACADEMICA, nombre_archivo="Búsqueda binaria",
                prefijo="w3")
    # La segunda estudiante busca en su propio objetivo de Recursividad.
    from app.security import create_token
    with Session(engine) as db:
        from app.models import Estudiante
        ana = db.exec(select(Estudiante).where(Estudiante.id_google == "otra-persona")).one()
    client.cookies.set("session", create_token(ana.id_estudiante))
    assert client.get("/api/auth/me").json()["id"] != otro
    plan_ana = crear_plan(client)
    assert client.get(f"/api/plans/{plan_ana}/search", params={"q": "caso base"}).json() == []
    # Y no puede buscar en el objetivo del otro estudiante.
    assert client.get(f"/api/plans/{plan_otro}/search", params={"q": "caso base"}).status_code == 404


def test_reindexar_no_duplica(client):
    tema = id_tema("Recursividad")
    for _ in range(2):
        rag.indexar([rag.Seccion("texto académico", "")], id_material=0, id_tema=tema, fuente=rag.ACADEMICA,
                    nombre_archivo="Recursión", prefijo="w9")
    assert len(rag.almacen().get()["ids"]) == 1


def test_script_de_fuentes_indexa_todos_los_temas(client, tmp_path, monkeypatch):
    from app import fuentes
    monkeypatch.setattr(fuentes, "CARPETA", tmp_path)
    monkeypatch.setattr(fuentes, "descargar", lambda t: (f"Intro de {t}.\n== Definición ==\nTexto de {t}.",
                                                         f"https://es.wikipedia.org/wiki/{t}"))
    fuentes.main()
    metas = rag.almacen().get()["metadatas"]
    assert {m["fuente"] for m in metas} == {"académica"} and {m["id_material"] for m in metas} == {0}
    assert len({m["id_tema"] for m in metas}) == 4
    assert "Definición" in {m["pagina_o_seccion"] for m in metas}
    assert (tmp_path / "recursividad" / "pila_de_llamadas.txt").read_text(encoding="utf-8").startswith("https://")


def test_descartar_quita_referencias():
    texto = "Intro\n== Uso ==\nalgo\n== Referencias ==\n1. Libro"
    assert fuentes_corte(texto) == "Intro\n== Uso ==\nalgo"


def fuentes_corte(texto):
    from app.fuentes import DESCARTAR
    m = DESCARTAR.search(texto)
    return texto[:m.start()].strip()


def test_el_servidor_retoma_el_material_pendiente_al_arrancar(client):
    """Simula un reinicio a mitad del procesamiento: el material queda "pendiente" sin que nadie lo suba de nuevo."""
    from app.main import _reprocesar_pendientes
    client.post("/api/auth/dev")
    plan = crear_plan(client)
    subir(client, plan, "apunte.md", "# Caso base\nTexto de prueba.".encode())
    with Session(engine) as db:
        mat = db.exec(select(MaterialEstudio)).first()
        id_material = mat.id_material
        mat.estado = "pendiente"  # como si el reinicio hubiera cortado el procesamiento acá
        db.add(mat); db.commit()
    import asyncio
    asyncio.run(_reprocesar_pendientes())
    with Session(engine) as db:
        assert db.get(MaterialEstudio, id_material).estado == "procesado"
